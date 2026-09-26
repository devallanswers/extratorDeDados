#!/usr/bin/env python3
"""
Extrator de dados de fichas de inscrição

Uso:
    python extrator_de_dados.py                   # processa pasta 'fichas/'
    python extrator_de_dados.py imagem1.jpg        # arquivo específico
    python extrator_de_dados.py pasta/             # outra pasta

- Usa Google Gemini para OCR de caligrafia (gratuito, 1500 req/dia)
- Recorta a foto 3x4 automaticamente da própria ficha
- Salva foto 3x4 na planilha Excel
- Incremental: não duplica fichas já processadas
"""

import json
import os
import re
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import google.generativeai as genai
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

# ──────────────────────────────────────────────────────────────────────────────
# CONFIGURAÇÕES
# ──────────────────────────────────────────────────────────────────────────────

PASTA_FICHAS  = "fichas"
ARQUIVO_SAIDA = "dados_extraidos.xlsx"
LOG_ERROS     = "erros_processamento.txt"

# Posição proporcional da caixinha de foto 3x4 na ficha
FOTO_LEFT   = 0.762
FOTO_TOP    = 0.150
FOTO_RIGHT  = 0.956
FOTO_BOTTOM = 0.250

CAMPOS = [
    ("Foto 3x4",                 "_foto"),
    ("Nº Ficha",                 "no_ficha"),
    ("Nome",                     "nome"),
    ("Apelido",                  "apelido"),
    ("Data de Nascimento",       "data_nascimento"),
    ("Idade",                    "idade"),
    ("Endereço",                 "endereco"),
    ("Bairro",                   "bairro"),
    ("Ponto de Referência",      "ponto_referencia"),
    ("Telefone Adolescente",     "telefone_adolescente"),
    ("Estuda?",                  "estuda"),
    ("Onde estuda",              "onde_estuda"),
    ("Responsável 1 / Mãe",     "responsavel_mae"),
    ("Telefone Mãe",             "telefone_mae"),
    ("Responsável 2 / Pai",     "responsavel_pai"),
    ("Telefone Pai",             "telefone_pai"),
    ("Com quem mora",            "com_quem_mora"),
    ("Tem irmãos?",              "tem_irmaos"),
    ("Nomes dos irmãos",         "nomes_irmaos"),
    ("Transporte",               "transporte"),
    ("Religião da família",      "religiao"),
    ("É batizado?",              "batizado"),
    ("Fez 1ª Eucaristia?",       "primeira_eucaristia"),
    ("Onde fez 1ª Eucaristia",   "onde_eucaristia"),
    ("Vai às missas?",           "vai_missas"),
    ("Onde vai à missa",         "onde_missa"),
    ("Diagnóstico / Laudo",      "diagnostico"),
    ("Cuidado especial",         "cuidado_especial"),
    ("Restrição p/ dormir fora", "restricao_dormir"),
    ("Alergias",                 "alergias"),
    ("Usa medicamento diário?",  "medicamento"),
    ("Qual medicamento",         "qual_medicamento"),
    ("Tratamento de saúde",      "tratamento_saude"),
    ("Por que quer participar",  "motivacao"),
    ("Arquivo",                  "_arquivo"),
    ("Processado em",            "_data_proc"),
    ("Status",                   "_status"),
]

PROMPT = """Esta imagem é uma ficha de inscrição do EAC preenchida à mão.
Leia cada campo escrito na ficha e retorne um JSON com os valores reais escritos.

IMPORTANTE: copie exatamente o que está escrito. NÃO use frases descritivas.

Exemplo:
{
  "no_ficha": "007",
  "nome": "Maria Eduarda Silva Santos",
  "apelido": "Duda"
}

Preencha TODOS os campos abaixo com o que está escrito na ficha (null se em branco):

{
  "no_ficha": null,
  "nome": null,
  "apelido": null,
  "data_nascimento": null,
  "idade": null,
  "endereco": null,
  "bairro": null,
  "ponto_referencia": null,
  "telefone_adolescente": null,
  "estuda": null,
  "onde_estuda": null,
  "responsavel_mae": null,
  "telefone_mae": null,
  "responsavel_pai": null,
  "telefone_pai": null,
  "com_quem_mora": null,
  "tem_irmaos": null,
  "nomes_irmaos": null,
  "transporte": null,
  "religiao": null,
  "batizado": null,
  "primeira_eucaristia": null,
  "onde_eucaristia": null,
  "vai_missas": null,
  "onde_missa": null,
  "diagnostico": null,
  "cuidado_especial": null,
  "restricao_dormir": null,
  "alergias": null,
  "medicamento": null,
  "qual_medicamento": null,
  "tratamento_saude": null,
  "motivacao": null
}

Retorne SOMENTE o JSON, sem explicações, sem markdown, sem ```."""


# ──────────────────────────────────────────────────────────────────────────────
# OCR COM GEMINI
# ──────────────────────────────────────────────────────────────────────────────

def extrair_dados_gemini(caminho: Path, modelo) -> dict:
    """Envia imagem pro Gemini e extrai os dados da ficha."""
    img = PILImage.open(caminho)
    resposta = modelo.generate_content([PROMPT, img])
    texto = resposta.text.strip()

    # Remove markdown se vier
    texto = re.sub(r"^```(?:json)?\s*", "", texto)
    texto = re.sub(r"\s*```\s*$", "", texto)

    # Pega só o JSON se vier texto antes/depois
    m = re.search(r"\{.*\}", texto, re.DOTALL)
    if m:
        texto = m.group(0)

    return json.loads(texto)


# ──────────────────────────────────────────────────────────────────────────────
# RECORTE DA FOTO 3X4
# ──────────────────────────────────────────────────────────────────────────────

def recortar_foto_3x4(caminho_ficha: Path) -> Path | None:
    """Recorta a região da foto 3x4 da ficha e salva num arquivo temporário."""
    try:
        img = PILImage.open(caminho_ficha)
        w, h = img.size
        foto = img.crop((
            int(w * FOTO_LEFT),
            int(h * FOTO_TOP),
            int(w * FOTO_RIGHT),
            int(h * FOTO_BOTTOM),
        ))
        tmp = Path(tempfile.mktemp(suffix=".png"))
        foto.save(tmp)
        return tmp
    except Exception as e:
        print(f"  ⚠️  Não foi possível recortar foto 3x4: {e}")
        return None


# ──────────────────────────────────────────────────────────────────────────────
# PLANILHA EXCEL
# ──────────────────────────────────────────────────────────────────────────────

COR_HEADER = "1B3A6B"
COR_OK     = "E8F5E9"
COR_WARN   = "FFF9C4"
COR_ERRO   = "FFEBEE"
LINHA_ALTURA = 85


def _borda():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def criar_planilha(caminho: str) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados extraídos"
    ws.row_dimensions[1].height = 38
    for col, (header, chave) in enumerate(CAMPOS, 1):
        c = ws.cell(row=1, column=col, value=header)
        c.font      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill      = PatternFill("solid", fgColor=COR_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = _borda()
        ws.column_dimensions[get_column_letter(col)].width = 12 if chave == "_foto" else max(14, len(header) + 2)
    wb.save(caminho)
    return wb


def abrir_ou_criar(caminho: str):
    if os.path.exists(caminho):
        wb = openpyxl.load_workbook(caminho)
        ws = wb.active
        print(f"  → Planilha existente: {caminho} ({ws.max_row - 1} fichas)")
    else:
        wb = criar_planilha(caminho)
        ws = wb.active
        print(f"  → Nova planilha criada: {caminho}")
    return wb, ws


def ja_processados(ws) -> set:
    col = next((i + 1 for i, (_, k) in enumerate(CAMPOS) if k == "_arquivo"), None)
    if not col or ws.max_row < 2:
        return set()
    return {ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}


def adicionar_linha(ws, dados: dict, arquivo: str, status: str, foto_tmp: Path | None):
    linha = ws.max_row + 1
    cor  = COR_OK if status == "OK" else (COR_WARN if status == "ATENÇÃO" else COR_ERRO)
    fill = PatternFill("solid", fgColor=cor)
    al   = Alignment(vertical="center", wrap_text=True)
    dados["_arquivo"]   = arquivo
    dados["_data_proc"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    dados["_status"]    = status
    col_foto = None
    for col, (_, chave) in enumerate(CAMPOS, 1):
        if chave == "_foto":
            col_foto = col
            c = ws.cell(row=linha, column=col, value="")
            c.fill = fill; c.border = _borda()
            continue
        val = dados.get(chave) or ""
        c = ws.cell(row=linha, column=col, value=str(val) if val else "")
        c.alignment = al; c.fill = fill; c.border = _borda()
    ws.row_dimensions[linha].height = LINHA_ALTURA
    if col_foto and foto_tmp and foto_tmp.exists():
        try:
            xl_img = XLImage(str(foto_tmp))
            xl_img.width = 90; xl_img.height = 110
            ws.add_image(xl_img, f"{get_column_letter(col_foto)}{linha}")
        except Exception as e:
            print(f"  ⚠️  Erro ao inserir foto: {e}")


def log_erro(arquivo: str, erro: str):
    with open(LOG_ERROS, "a", encoding="utf-8") as f:
        f.write(f"[{datetime.now().strftime('%d/%m/%Y %H:%M')}] {arquivo}: {erro}\n")


# ──────────────────────────────────────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────────────────────────────────────

def coletar_arquivos(entrada) -> list[Path]:
    entrada = Path(entrada)
    exts = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
    if entrada.is_file():
        return [entrada]
    elif entrada.is_dir():
        return sorted(f for f in entrada.iterdir() if f.suffix.lower() in exts)
    else:
        print(f"ERRO: '{entrada}' não existe.")
        sys.exit(1)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    entrada = sys.argv[1] if len(sys.argv) > 1 else PASTA_FICHAS

    print("=" * 58)
    print("  EXTRATOR DE DADOS — Google Gemini")
    print("=" * 58)

    # Configura Gemini
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        print("\n❌ Chave do Gemini não configurada!")
        print("   Configure a variável de ambiente GEMINI_API_KEY antes de executar.")
        sys.exit(1)

    genai.configure(api_key=api_key)
    modelo = genai.GenerativeModel("gemini-1.5-flash")
    print(f"\n  Google Gemini ✅")

    arquivos = coletar_arquivos(entrada)
    if not arquivos:
        print(f"\nNenhum arquivo encontrado em: {entrada}")
        sys.exit(0)

    print(f"\n📋 {len(arquivos)} ficha(s) para processar\n")

    wb, ws = abrir_ou_criar(ARQUIVO_SAIDA)
    processados_antes = ja_processados(ws)
    fotos_tmp = []
    ok = erros = pulados = 0

    for idx, caminho in enumerate(arquivos, 1):
        id_arquivo = caminho.name
        print(f"[{idx}/{len(arquivos)}] {id_arquivo}")

        if id_arquivo in processados_antes:
            print("  ✓ Já processado, pulando.\n")
            pulados += 1
            continue

        try:
            foto_tmp = recortar_foto_3x4(caminho)
            if foto_tmp:
                fotos_tmp.append(foto_tmp)
                print("  → Foto 3x4 recortada ✅")

            print("  → Enviando para Gemini...")
            dados = extrair_dados_gemini(caminho, modelo)

            importantes = ["nome", "data_nascimento", "telefone_adolescente"]
            vazios = [c for c in importantes if not dados.get(c)]
            status = "ATENÇÃO" if vazios else "OK"
            if vazios:
                print(f"  ⚠️  Campos vazios: {', '.join(vazios)}")

            adicionar_linha(ws, dados, id_arquivo, status, foto_tmp)
            wb.save(ARQUIVO_SAIDA)
            ok += 1
            print(f"  ✅ {dados.get('nome') or '(sem nome)'}\n")

        except json.JSONDecodeError as e:
            print(f"  ❌ ERRO ao interpretar resposta: {e}\n")
            adicionar_linha(ws, {}, id_arquivo, "ERRO", None)
            wb.save(ARQUIVO_SAIDA)
            log_erro(id_arquivo, f"JSON inválido: {e}")
            erros += 1

        except Exception as e:
            print(f"  ❌ ERRO: {e}\n")
            adicionar_linha(ws, {}, id_arquivo, "ERRO", None)
            wb.save(ARQUIVO_SAIDA)
            log_erro(id_arquivo, str(e))
            erros += 1

    for f in fotos_tmp:
        try: f.unlink()
        except: pass

    print("=" * 58)
    print(f"  ✅ Processadas : {ok}")
    print(f"  ⏭️  Puladas     : {pulados}")
    print(f"  ❌ Erros       : {erros}")
    print(f"  📊 Planilha    : {ARQUIVO_SAIDA}")
    if erros:
        print(f"  📝 Log erros   : {LOG_ERROS}")
    print("=" * 58)


if __name__ == "__main__":
    main()